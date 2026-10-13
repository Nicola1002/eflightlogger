"""Storage backends: a local xlsx file (for testing) or a Google Sheet via the Sheets API (OAuth).

Both expose the same interface:
    read_table(sheet)                      -> list[dict]   (header-name -> value)
    append_row(sheet, row)                 -> None         (values placed by header name,
                                                            unknown headers are added at the end)
    ensure_sheet(sheet, headers, rows)     -> None         (create sheet if missing)

Row values may be str, int, float, None, datetime.date or datetime.time.
"""
import datetime as dt
import shutil
import threading
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

HEADER_SCAN_ROWS = 15


def _clean(v):
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    return v


class LocalXlsxBackend:
    def __init__(self, path, seed=None, base_dir=Path(".")):
        self.path = (base_dir / path).resolve()
        self.lock = threading.Lock()
        if not self.path.exists():
            self.path.parent.mkdir(parents=True, exist_ok=True)
            if seed and (base_dir / seed).exists():
                shutil.copy(base_dir / seed, self.path)
            else:
                from openpyxl import Workbook
                Workbook().save(self.path)

    @staticmethod
    def _header(ws):
        for r in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
            cells = [(c.column, str(c.value).strip()) for c in ws[r] if c.value not in (None, "")]
            if len(cells) >= 2:
                return r, dict((name, col) for col, name in cells)
        return None, {}

    def read_table(self, sheet):
        with self.lock:
            wb = load_workbook(self.path, data_only=True)
            if sheet not in wb.sheetnames:
                return []
            ws = wb[sheet]
            hrow, cols = self._header(ws)
            rows = []
            for r in range(hrow + 1, ws.max_row + 1) if hrow else []:
                row = {n: _clean(ws.cell(r, c).value) for n, c in cols.items()}
                if any(v not in (None, "") for v in row.values()):
                    rows.append(row)
            return rows

    def ensure_sheet(self, sheet, headers, rows=()):
        with self.lock:
            wb = load_workbook(self.path)
            if sheet in wb.sheetnames:
                return
            ws = wb.create_sheet(sheet)
            ws.append(list(headers))
            for r in rows:
                ws.append(list(r))
            wb.save(self.path)

    def append_row(self, sheet, row):
        with self.lock:
            wb = load_workbook(self.path)
            if sheet not in wb.sheetnames:
                raise RuntimeError(f"Blatt '{sheet}' fehlt in {self.path.name}")
            ws = wb[sheet]
            hrow, cols = self._header(ws)
            if not hrow:
                raise RuntimeError(f"Blatt '{sheet}' hat keine Kopfzeile")
            # add headers we do not know yet, right of the last one
            for name in row:
                if name not in cols:
                    col = max(cols.values()) + 1
                    head = ws.cell(hrow, col, name)
                    prev = ws.cell(hrow, col - 1)
                    if prev.has_style:
                        head._style = copy(prev._style)
                    cols[name] = col
            # first free row below the last row that has data in a header column
            last = hrow
            for r in range(ws.max_row, hrow, -1):
                if any(ws.cell(r, c).value not in (None, "") for c in cols.values()):
                    last = r
                    break
            new = last + 1
            for name, col in cols.items():
                src = ws.cell(last, col)
                dst = ws.cell(new, col)
                if last > hrow and src.has_style:
                    dst._style = copy(src._style)
                if name not in row:
                    continue
                v = row[name]
                dst.value = v
                if isinstance(v, dt.datetime):
                    dst.number_format = "DD.MM.YYYY HH:MM"
                elif isinstance(v, dt.date):
                    dst.number_format = "DD.MM.YYYY"
                elif isinstance(v, dt.time):
                    dst.number_format = "HH:MM"
            wb.save(self.path)


class NotAuthenticated(Exception):
    """No (valid) Google login yet – the user has to sign in first."""


SHEETS_EPOCH = dt.datetime(1899, 12, 30)
# Sheets number-format patterns (lower case, unlike Excel)
FMT_DATETIME = {"type": "DATE_TIME", "pattern": "dd.mm.yyyy hh:mm"}
FMT_DATE = {"type": "DATE", "pattern": "dd.mm.yyyy"}
FMT_TIME = {"type": "TIME", "pattern": "hh:mm"}


def _a1(sheet):
    return "'" + sheet.replace("'", "''") + "'"


def _from_serial(n, kind):
    d = SHEETS_EPOCH + dt.timedelta(seconds=round(n * 86400))
    return d.time().isoformat() if kind == "TIME" else d.isoformat()


def _cell_value(cell):
    """Sheets API cell -> plain value; dates/times become ISO strings like in the xlsx backend."""
    ev = cell.get("effectiveValue") or {}
    if "numberValue" in ev:
        n = ev["numberValue"]
        kind = ((cell.get("effectiveFormat") or {}).get("numberFormat") or {}).get("type")
        if kind in ("DATE", "DATE_TIME", "TIME"):
            return _from_serial(n, kind)
        return n
    if "stringValue" in ev:
        return ev["stringValue"]
    if "boolValue" in ev:
        return ev["boolValue"]
    return None


def _find_header(grid):
    """First row among the first HEADER_SCAN_ROWS with >= 2 filled cells -> (row index, {name: col index})."""
    for r, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        cells = [(c, str(v).strip()) for c, v in enumerate(row) if v not in (None, "")]
        if len(cells) >= 2:
            return r, {name: c for c, name in cells}
    return None, {}


def _write_value(v):
    """Python value -> (Sheets ExtendedValue, number format or None)."""
    if isinstance(v, bool):
        return {"boolValue": v}, None
    if isinstance(v, dt.datetime):
        return {"numberValue": (v - SHEETS_EPOCH).total_seconds() / 86400}, FMT_DATETIME
    if isinstance(v, dt.date):
        return {"numberValue": (dt.datetime(v.year, v.month, v.day) - SHEETS_EPOCH).days}, FMT_DATE
    if isinstance(v, dt.time):
        secs = v.hour * 3600 + v.minute * 60 + v.second
        return {"numberValue": secs / 86400}, FMT_TIME
    if isinstance(v, (int, float)):
        return {"numberValue": v}, None
    return {"stringValue": str(v)}, None  # stringValue is literal text, never a formula


class GoogleSheetsApiBackend:
    """Google Sheets API v4 with the signed-in user's OAuth credentials (see auth.py)."""

    def __init__(self, spreadsheet_id, get_credentials):
        if not spreadsheet_id:
            raise RuntimeError("config: google.spreadsheet_id muss gesetzt sein")
        self.id = spreadsheet_id
        self.get_credentials = get_credentials
        self.lock = threading.Lock()

    def _svc(self):
        from googleapiclient.discovery import build
        creds = self.get_credentials()
        if creds is None:
            raise NotAuthenticated("Nicht bei Google angemeldet")
        # httplib2 is not thread-safe -> a fresh client per call
        return build("sheets", "v4", credentials=creds, cache_discovery=False).spreadsheets()

    def _meta(self, svc):
        r = svc.get(spreadsheetId=self.id, fields="sheets.properties").execute(num_retries=2)
        return {s["properties"]["title"]: s["properties"] for s in r.get("sheets", [])}

    def _grid(self, svc, sheet):
        r = svc.get(spreadsheetId=self.id, ranges=[_a1(sheet)], includeGridData=True,
                    fields="sheets.data.rowData.values(effectiveValue,effectiveFormat.numberFormat.type)"
                    ).execute(num_retries=2)
        data = r["sheets"][0]["data"][0] if r.get("sheets") and r["sheets"][0].get("data") else {}
        return [[_cell_value(c) for c in row.get("values", [])] for row in data.get("rowData", [])]

    def read_table(self, sheet):
        svc = self._svc()
        if sheet not in self._meta(svc):
            return []
        grid = self._grid(svc, sheet)
        hrow, cols = _find_header(grid)
        rows = []
        for raw in grid[hrow + 1:] if hrow is not None else []:
            row = {n: (raw[c] if c < len(raw) else None) for n, c in cols.items()}
            if any(v not in (None, "") for v in row.values()):
                rows.append(row)
        return rows

    def ensure_sheet(self, sheet, headers, rows=()):
        with self.lock:
            svc = self._svc()
            if sheet in self._meta(svc):
                return
            svc.batchUpdate(spreadsheetId=self.id, body={"requests": [
                {"addSheet": {"properties": {"title": sheet}}}]}).execute(num_retries=2)
            svc.values().update(spreadsheetId=self.id, range=f"{_a1(sheet)}!A1", valueInputOption="RAW",
                                body={"values": [list(headers)] + [list(r) for r in rows]}
                                ).execute(num_retries=2)

    def append_row(self, sheet, row):
        with self.lock:
            svc = self._svc()
            props = self._meta(svc).get(sheet)
            if props is None:
                raise RuntimeError(f"Blatt '{sheet}' fehlt in der Google-Tabelle")
            grid = self._grid(svc, sheet)
            hrow, cols = _find_header(grid)
            if hrow is None:
                raise RuntimeError(f"Blatt '{sheet}' hat keine Kopfzeile")
            sid = props["sheetId"]
            reqs, new_heads = [], {}
            # add headers we do not know yet, right of the last one
            for name in row:
                if name not in cols:
                    col = max(cols.values()) + 1
                    new_heads[name] = cols[name] = col
            # first free row below the last row that has data in a header column
            last = hrow
            for r in range(len(grid) - 1, hrow, -1):
                if any(c < len(grid[r]) and grid[r][c] not in (None, "") for c in cols.values()):
                    last = r
                    break
            new = last + 1

            grow = props["gridProperties"]
            if new >= grow["rowCount"]:
                reqs.append({"appendDimension": {"sheetId": sid, "dimension": "ROWS",
                                                 "length": new + 1 - grow["rowCount"]}})
            need_cols = max(cols.values()) + 1
            if need_cols > grow["columnCount"]:
                reqs.append({"appendDimension": {"sheetId": sid, "dimension": "COLUMNS",
                                                 "length": need_cols - grow["columnCount"]}})

            def rng(r, c):
                return {"sheetId": sid, "startRowIndex": r, "endRowIndex": r + 1,
                        "startColumnIndex": c, "endColumnIndex": c + 1}

            def put(r, c, v):
                ev, fmt = _write_value(v)
                cell, fields = {"userEnteredValue": ev}, "userEnteredValue"
                if fmt:
                    cell["userEnteredFormat"] = {"numberFormat": fmt}
                    fields += ",userEnteredFormat.numberFormat"
                return {"updateCells": {"rows": [{"values": [cell]}], "fields": fields,
                                        "start": {"sheetId": sid, "rowIndex": r, "columnIndex": c}}}

            for name, col in new_heads.items():  # new header cells inherit the style of their left neighbour
                reqs.append({"copyPaste": {"source": rng(hrow, col - 1), "destination": rng(hrow, col),
                                           "pasteType": "PASTE_FORMAT"}})
                reqs.append(put(hrow, col, name))
            if last > hrow:  # copy the previous row's formatting, then set values
                reqs.append({"copyPaste": {
                    "source": {"sheetId": sid, "startRowIndex": last, "endRowIndex": last + 1,
                               "startColumnIndex": 0, "endColumnIndex": need_cols},
                    "destination": {"sheetId": sid, "startRowIndex": new, "endRowIndex": new + 1,
                                    "startColumnIndex": 0, "endColumnIndex": need_cols},
                    "pasteType": "PASTE_FORMAT"}})
            for name, v in row.items():
                if v is not None:
                    reqs.append(put(new, cols[name], v))
            svc.batchUpdate(spreadsheetId=self.id, body={"requests": reqs}).execute(num_retries=2)


def make_backend(cfg, base_dir, get_credentials=None):
    if cfg["backend"] == "google":
        return GoogleSheetsApiBackend(cfg["google"]["spreadsheet_id"], get_credentials)
    return LocalXlsxBackend(cfg["local"]["path"], cfg["local"].get("seed"), base_dir)
